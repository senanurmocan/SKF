"""Exercise browser-staged synthetic Excel inputs through the invoice pipeline.

The fixtures deliberately contain no operational invoice data and do not depend
on private monthly workbooks. Expected totals are independent business examples.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

import app
from core.upload_inputs import UploadWorkspace


class UploadedWorkbook(io.BytesIO):
    def __init__(self, name: str, data: bytes) -> None:
        super().__init__(data)
        self.name = name
        self.size = len(data)


def excel_bytes(headers, rows, *, leading_rows=()):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Örnek Veriler"
    for row in leading_rows:
        worksheet.append(row)
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


# Keep these fields aligned with the configurable mapping contract, rather than
# using actual regional invoices or fixed operational column positions.
REGION_FIELDS = [
    ("ETSO", "ETSO Kodu"),
    ("Müşteri", "Ad Soyad"),
    ("Tarife Grubu", "Fatura Abone Grubu2"),
    ("AG-OG", "AG/OG"),
    ("Terim", "Terim"),
    ("Kurulu Güç", "Kurulu Güç"),
    ("Sözleşme Gücü", "Sözleşme Gücü"),
    ("Aktif Enerji (kwh)", "Toplam Kwh"),
    ("Aktif Enerji TL", "Toplam enerji Bedeli(Aktif+Trafo Kaybı+Ek Tüketim+PSH)"),
    ("Dağıtım Bedeli", "Toplam Dağıtım Bedeli"),
    ("Güç Bedeli", "Güç Bedeli"),
    ("Güç aşım Bedeli", "Güç aşım Bedeli"),
    ("Reaktif Bedel", "Reaktif Bedel + Enerji Bedeli Net Reaktif İhlal"),
    ("BTV", "Btv Bedeli"),
    ("Fiyat Farkı", "Fiyat Farkı"),
    ("Kdv", "Kdv"),
    ("KDV Matrahı", "KDV Matrahı"),
    ("TOPLAM", "Fatura Tutarı"),
]
REGION_HEADERS = [source for _, source in REGION_FIELDS] + [
    "Hizmet Noktası No", "Dağıtım Hizmet Noktası No", "Fatura Abone Grubu",
    "Fatura Segmenti  Muhasebe İptal Tarihi", "Dağıtım Bölgesi",
]
MANAGEMENT_HEADERS = [
    "Dağıtım Bölgesi", "ETSO", "Müşteri", "Tarife Grubu", "Abone Grubu",
    "Aktif Enerji Tüketim (kWh)", "Dağıtım Bedeli (TL)", "TL Güç",
    "Güç Aşım Bedeli (TL)", "Reaktif Bedel (TL)", "Tük Güç", "Durum",
]
FIRST_ETSO = "40XTEST000000001"
SECOND_ETSO = "40XTEST000000002"


def region_row(etso=FIRST_ETSO, **overrides):
    values = dict(zip(REGION_HEADERS, [
        etso, "Sentetik Müşteri", "Kamu/Özel Sektör ve Diğer", "AG", "Tek Terimli",
        20, 10, 100, 200, 30, 4, 5, 6, 7, 8, 52, 260, 312,
        "HN-1", "DHN-1", "Yalnız tanıma için", None, "Akdeniz EDAŞ",
    ]))
    values.update(overrides)
    return [values.get(header) for header in REGION_HEADERS]


def management_rows():
    return [
        ["Akdeniz EDAŞ", FIRST_ETSO, "Sentetik A", "DSK Ticarethane TT AG",
         "Ticarethane AG", 200, 99, 0, 0, 0, 10, "Kesin"],
        ["Akdeniz EDAŞ", SECOND_ETSO, "Sentetik B", "DSK Sanayi ÇT OG",
         "Sanayi OG", 500, 888, 45, 2, 3, 10, "Kesin"],
        ["Akdeniz EDAŞ", "40XTESTDRAFT", "Sentetik Taslak", "DSK Ticarethane TT AG",
         "Ticarethane AG", 999, 99, 0, 0, 0, 10, "Taslak"],
    ]


@pytest.fixture
def staged_inputs():
    """Provide only synthetic mapping/rate workbooks in a browser workspace."""
    mapping = Workbook()
    tariffs = mapping.active
    tariffs.title = "Tarifeler"
    tariffs.append([
        "Tarife Grubu", "Dağıtım Sistemi Kullanıcıları", "Abone / Kullanıcı Grubu",
        "Birim", "Sütun", "Satır",
    ])
    tariffs.append(["DSK Ticarethane TT AG", "AG Tek Terim", "Ticarethane", "kr/kWh", "I", 18])
    tariffs.append(["DSK Sanayi ÇT OG", "OG Çift Terim", "Sanayi", "kr/kWh", "I", 7])
    layout = mapping.create_sheet("Bizim Kestiğimiz 2")
    layout.append([None, "Tarife Grubu", "Terim Tipi", "Bağlantı Tipi"])
    layout.append([None, "DSK Ticarethane TT AG", "Tek Terim", "AG"])
    layout.append([None, "DSK Sanayi ÇT OG", "Çift Terim", "OG"])
    region = mapping.create_sheet("Bölge Bizim Kestiğimiz 2")
    field_headers, source_headers = [], []
    for index, (field, source) in enumerate(REGION_FIELDS, start=1):
        field_headers.extend([field, "Sütun"])
        source_headers.extend([source, get_column_letter(index)])
    region.append(field_headers)
    region.append(source_headers)
    region.append(["Bizim Kestiğimiz 2'ye Alırken Düzeltme"])
    region.append(["Başlık", "Yükelenen", "Düzeltme"])
    region.append(["Tarife Grubu", "Kamu/Özel Sektör ve Diğer", "Ticarethane"])
    region.append(["Terim", "Tek Terimli", "Tek Terim"])
    region.append(["Terim", "Çift Terimli", "Çift Terim"])
    buffer = io.BytesIO()
    mapping.save(buffer)
    mapping.close()

    rates = Workbook()
    rate_sheet = rates.active
    rate_sheet.cell(7, 3, "OG Çift Terim")
    rate_sheet.cell(7, 4, "Sanayi")
    rate_sheet.cell(7, 9, 125)
    rate_sheet.cell(7, 10, 450)
    rate_sheet.cell(18, 3, "AG Tek Terim")
    rate_sheet.cell(18, 4, "Ticarethane")
    rate_sheet.cell(18, 9, 32.5)
    rate_buffer = io.BytesIO()
    rates.save(rate_buffer)
    rates.close()

    with UploadWorkspace() as workspace:
        mapping_path = workspace.stage_file(UploadedWorkbook("başlıklar.xlsx", buffer.getvalue()), "mapping")
        rates_path = workspace.stage_file(UploadedWorkbook("tarifeler.xlsx", rate_buffer.getvalue()), "rates")
        yield workspace, mapping_path, rates_path


def run_invoice(staged_inputs, filename, data, *, reconciliation_data=None):
    workspace, mapping_path, rates_path = staged_inputs
    invoice_path = workspace.stage_file(UploadedWorkbook(filename, data), "invoice")
    reconciliation_path = None
    if reconciliation_data is not None:
        reconciliation_path = workspace.stage_file(
            UploadedWorkbook("adı tamamen farklı sayaç.xlsx", reconciliation_data), "reconciliation"
        )
    result = app.run_analysis_pipeline(
        root_folder="", mapping_path=mapping_path, status_callback=None,
        invoice_file=invoice_path, system_usage_file=rates_path,
        reconciliation_consumption_file=reconciliation_path,
    )
    _, stats, _, output, _, sources = result
    assert stats["invoice_only_mode"] is True
    assert stats["invoice_file_count"] == 1
    assert stats["invoice_file_errors"] == []
    assert len(sources) == 1
    return result, load_workbook(io.BytesIO(output), data_only=False)


def records(worksheet):
    values = worksheet.iter_rows(values_only=True)
    headers = next(values)
    return [dict(zip(headers, row)) for row in values]


@pytest.mark.parametrize("filename", ["rasgele ad.xlsx", "Bölge İcmal diye adlandırılmış.xlsx"])
def test_original_invoice_management_survives_browser_upload_and_renaming(staged_inputs, filename):
    result, workbook = run_invoice(
        staged_inputs, filename, excel_bytes(MANAGEMENT_HEADERS, management_rows())
    )
    try:
        preview, _, _, _, _, sources = result
        assert sources[0]["invoice_format"] == "invoice_management"
        assert len(preview) == 3
        assert sources[0]["calculated_distribution_amount_count"] == 3
        first, second, draft = records(workbook["Bizim Kestiğimiz 2"])
        assert first["Etso"] == FIRST_ETSO
        assert first["Abone Grubu"] == "Ticarethane"
        assert first["Terim Tipi"] == "Tek Terim"
        assert first["Bağlantı Tipi"] == "AG"
        assert first["Dağıtım Bedeli (TL)"] == pytest.approx(65)
        assert first["Durum"] == "Kesin"
        assert second["Abone Grubu"] == "Sanayi"
        assert second["Dağıtım Bedeli (TL)"] == pytest.approx(625)
        assert second["Güç Bedeli(TL)"] == 45
        assert second["Terim Tipi"] == "Çift Terim"
        assert second["Güç kWh"].startswith("=IFERROR(")
        assert draft["Durum"] == "Taslak"
        assert "Güç Fark" in first
        assert "Karşılaştırma" not in workbook.sheetnames
    finally:
        workbook.close()


@pytest.mark.parametrize("filename", ["tamamen farklı.xlsx", "FaturaYönetimi.xlsx"])
def test_region_upload_aggregates_only_uncancelled_rows_with_mapping_corrections(staged_inputs, filename):
    invoice_rows = [
        region_row(),
        region_row(**{"Kurulu Güç": 15, "Sözleşme Gücü": 12, "Toplam Kwh": 150,
                      "Fatura Segmenti  Muhasebe İptal Tarihi": "   "}),
        region_row(**{"Toplam Kwh": 9000, "Kurulu Güç": 999, "Sözleşme Gücü": 999,
                      "Fatura Segmenti  Muhasebe İptal Tarihi": "2026-10-01"}),
        region_row(SECOND_ETSO, **{"Terim": "Çift Terimli", "Toplam Kwh": 300}),
        region_row("40XCANCELLEDONLY", **{"Fatura Segmenti  Muhasebe İptal Tarihi": "iptal"}),
        region_row(None, **{"Toplam Kwh": 100000}),
    ]
    result, workbook = run_invoice(staged_inputs, filename, excel_bytes(REGION_HEADERS, invoice_rows))
    try:
        preview, _, _, _, logs, sources = result
        source = sources[0]
        assert source["invoice_format"] == "region_summary"
        assert source["region_invoice_source_row_count"] == 3
        assert source["region_invoice_etso_count"] == 2
        assert source["region_invoice_duplicate_row_count"] == 1
        assert len(preview) == 2
        first, second = records(workbook["Bizim Kestiğimiz 2"])
        assert first["Etso"] == FIRST_ETSO
        assert first["Tarife Grubu"] == "Ticarethane"
        assert first["Terim Tipi"] == "Tek Terim"
        assert second["Terim Tipi"] == "Çift Terim"
        assert first["Sözleşme Gücü (kW)"] == 12
        assert first["Kurulu Güç (kW)"] == 20
        assert first["Aktif Enerji Tüketim (kWh)"] == 250
        assert first["Aktif Enerji TL"] == 400
        assert first["Dağıtım Bedeli (TL)"] == 60
        assert first["BTV Bedeli (TL)"] == 14
        assert first["Fiyat Farkı (TL)"] == 16
        assert first["KDV"] == 104
        assert first["KDV Matrahı (TL)"] == 520
        assert first["Toplam (TL)"] == 624
        assert not {"Abone Grubu", "Güç Fark", "Durum"}.intersection(first)
        assert any("maksimum" in log for log in logs)
        assert "Karşılaştırma" not in workbook.sheetnames
    finally:
        workbook.close()


def reconciliation_bytes():
    return excel_bytes(
        ["Sayaç ETSO", "Çekiş Org", "Kayıplı Çekiş (MWh)"],
        [
            [FIRST_ETSO, "Evet", 0.1],
            ["40X-TEST000000001", " EVET ", 0.05],
            [FIRST_ETSO, "Hayır", 90],
            [SECOND_ETSO, "Evet", 0.5],
            ["40XEXTRA", "Evet", 1],
            [None, "Evet", 99],
        ],
        leading_rows=[["Tamamen sentetik sayaç verileri"], []],
    )


def test_browser_reconciliation_filters_draw_flags_aggregates_and_converts_to_kwh(staged_inputs):
    workspace, _, _ = staged_inputs
    path = workspace.stage_file(UploadedWorkbook("rasgele tüketim.xlsx", reconciliation_bytes()), "reconciliation")
    values, stats = app.load_reconciliation_consumption_by_etso(path)
    assert values == {FIRST_ETSO: pytest.approx(150), SECOND_ETSO: 500, "40XEXTRA": 1000}
    assert stats["reconciliation_consumption_source_row_count"] == 4
    assert stats["reconciliation_consumption_etso_count"] == 3


@pytest.mark.parametrize("invoice_kind", ["invoice_management", "region_summary"])
def test_optional_reconciliation_works_without_distribution_and_is_not_implicitly_reused(staged_inputs, invoice_kind):
    if invoice_kind == "invoice_management":
        headers, rows = MANAGEMENT_HEADERS, management_rows()
    else:
        headers, rows = REGION_HEADERS, [
            region_row(**{"Toplam Kwh": 200}),
            region_row(SECOND_ETSO, **{"Toplam Kwh": 500}),
        ]
    data = excel_bytes(headers, rows)
    result, workbook = run_invoice(
        staged_inputs, "seçilen fatura.xlsx", data, reconciliation_data=reconciliation_bytes()
    )
    try:
        source = result[5][0]
        assert source["reconciliation_consumption_matched_count"] == 2
        assert source["reconciliation_consumption_unmatched_count"] == 0
        comparison = workbook["Karşılaştırma"]
        first, second = records(comparison)
        assert first["ETSO"] == FIRST_ETSO
        assert first["Uzlaştırma Tüketim (kWh)"] == 150
        assert first["Aktif Enerji Tüketim (kWh) Dağıtım"] == "Dağıtım verisi yok"
        assert first["Dağıtım - Uzlaştırma Tüketim Farkı (kWh)"] == "-"
        assert second["Uzlaştırma Tüketim (kWh)"] == 500
        differences = source["our_vs_reconciliation_difference_rows"]
        assert len(differences) == 1
        assert differences[0]["ETSO"] == FIRST_ETSO
        assert differences[0]["Fark (kWh)"] == pytest.approx(50)
        assert source["distribution_vs_reconciliation_difference_rows"] == []
    finally:
        workbook.close()

    # Files may still exist in the per-session workspace; each analysis must
    # nevertheless require an explicit reconciliation selection.
    next_result, next_workbook = run_invoice(staged_inputs, "seçilen fatura.xlsx", data)
    try:
        assert "Karşılaştırma" not in next_workbook.sheetnames
        assert "our_vs_reconciliation_difference_rows" not in next_result[5][0]
    finally:
        next_workbook.close()


def test_reconciliation_reports_split_our_and_distribution_differences():
    workbook = Workbook()
    comparison = workbook.active
    comparison.title = "Karşılaştırma"
    comparison.append([
        "ETSO", "Dağıtım Şirketi", "Müşteri", "Tarife Grubu",
        "Aktif Enerji Tüketim (kWh)", "Aktif Enerji Tüketim (kWh) Dağıtım",
    ])
    comparison.append([FIRST_ETSO, "Akdeniz EDAŞ", "Sentetik A", "Ticarethane", 200, 150])
    comparison.append([SECOND_ETSO, "Akdeniz EDAŞ", "Sentetik B", "Sanayi", 500, 525])
    comparison.append(["40XMISSING", "Akdeniz EDAŞ", "Sentetik C", "Mesken", 400, 450])
    try:
        stats = app.add_reconciliation_consumption_columns(
            workbook, comparison, {FIRST_ETSO: 150, SECOND_ETSO: 500}
        )
        assert stats["reconciliation_consumption_matched_count"] == 2
        assert stats["reconciliation_consumption_unmatched_count"] == 1
        ours = stats["our_vs_reconciliation_difference_rows"]
        distribution = stats["distribution_vs_reconciliation_difference_rows"]
        assert [(r["ETSO"], r["Fark (kWh)"]) for r in ours] == [(FIRST_ETSO, 50)]
        assert [(r["ETSO"], r["Fark (kWh)"]) for r in distribution] == [(SECOND_ETSO, 25)]
        assert comparison.cell(4, 7).value == "-"
        assert comparison.auto_filter.ref == "A1:I4"
        for kind, rows, etso, difference in [
            ("ours", ours, FIRST_ETSO, 50),
            ("distribution", distribution, SECOND_ETSO, 25),
        ]:
            report = load_workbook(
                io.BytesIO(app.build_reconciliation_consumption_report(rows, kind)), data_only=True
            )
            try:
                exported = list(report["Fark Raporu"].values)
                assert len(exported) == 2
                assert exported[1][0] == etso
                assert exported[1][-1] == difference
                assert ("Bizim Kestiğimiz" if kind == "ours" else "Dağıtım") in exported[0][-1]
            finally:
                report.close()
    finally:
        workbook.close()
