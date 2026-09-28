#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Investigate the mismatch between extraction and reference"""

import openpyxl

# Load extracted data
wb = openpyxl.load_workbook('Çıkarılan_Veriler.xlsx')
ws = wb.active
rows = list(ws.iter_rows(values_only=True))

extracted_data = []
for row in rows[1:]:
    extracted_data.append({
        'Dağıtım Bölgesi': row[0],
        'Etso Kodu': row[1],
        'Müşteri': row[2],
        'Dağıtım Bedeli(TL)': row[9]
    })

# Filter positive only
extracted_filtered = [r for r in extracted_data if r['Dağıtım Bedeli(TL)'] >= 0]

# Create extraction set
def make_key(r):
    return (r['Etso Kodu'], r['Müşteri'] if r['Müşteri'] else '', round(r['Dağıtım Bedeli(TL)'], 2))

extracted_set = {make_key(r) for r in extracted_filtered}

# Load reference
ref_wb = openpyxl.load_workbook('Dağıtımın Kestiği Faturalar Özet.xlsx')
ref_ws = ref_wb['Dağıtımın Kestiği']
ref_rows = list(ref_ws.iter_rows(values_only=True))

reference_data = []
for row in ref_rows[1:]:
    reference_data.append({
        'Dağıtım Bölgesi': row[0],
        'Etso Kodu': row[1],
        'Müşteri': row[2],
        'Dağıtım Bedeli(TL)': row[9]
    })

reference_set = {make_key(r) for r in reference_data}

# Find missing and extra
missing = [r for r in reference_data if make_key(r) not in extracted_set]
extra = [r for r in extracted_filtered if make_key(r) not in reference_set]

print(f'Reference unique keys: {len(reference_set)}')
print(f'Extraction unique keys (positive): {len(extracted_set)}')
print(f'Missing from extraction: {len(missing)}')
print(f'Extra in extraction: {len(extra)}')

# Check for region distribution in missing
print("\n=== REGION DISTRIBUTION IN MISSING ===")
missing_by_region = {}
for r in missing:
    region = r['Dağıtım Bölgesi']
    missing_by_region[region] = missing_by_region.get(region, 0) + 1

for region, count in sorted(missing_by_region.items(), key=lambda x: -x[1]):
    print(f"  {region}: {count}")

# Check for region distribution in extra
print("\n=== REGION DISTRIBUTION IN EXTRA ===")
extra_by_region = {}
for r in extra:
    region = r['Dağıtım Bölgesi']
    extra_by_region[region] = extra_by_region.get(region, 0) + 1

for region, count in sorted(extra_by_region.items(), key=lambda x: -x[1]):
    print(f"  {region}: {count}")

# Check sample of missing records
print("\n=== SAMPLE OF MISSING RECORDS (first 10) ===")
for r in missing[:10]:
    print(f"  {r['Dağıtım Bölgesi']}: Etso={r['Etso Kodu']}, Müşteri={r['Müşteri']}, Bedel={r['Dağıtım Bedeli(TL)']}")

# Check sample of extra records
print("\n=== SAMPLE OF EXTRA RECORDS (first 10) ===")
for r in extra[:10]:
    print(f"  {r['Dağıtım Bölgesi']}: Etso={r['Etso Kodu']}, Müşteri={r['Müşteri']}, Bedel={r['Dağıtım Bedeli(TL)']}")

# Check if the issue is Etso Kodu format differences
print("\n=== ETSO KODU FORMAT ANALYSIS ===")
etso_types = {}
for r in extracted_filtered:
    etso = str(r['Etso Kodu'])
    if etso not in etso_types:
        etso_types[etso] = {'region': r['Dağıtım Bölgesi'], 'type': 'numeric' if etso.isdigit() else 'mixed'}
        
print(f"Sample Etso Kodu values:")
for i, (etso, info) in enumerate(etso_types.items()):
    if i < 20:
        print(f"  {etso}: {info['type']} ({info['region']})")
