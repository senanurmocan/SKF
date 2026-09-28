#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Discrepancy analysis after negative bedel filter"""

import openpyxl
from extract_and_compare import *

# Load extracted data
wb = openpyxl.load_workbook('Çıkarılan_Veriler.xlsx')
ws = wb.active
rows = list(ws.iter_rows(values_only=True))
extracted_data = []
for row in rows[1:]:
    extracted_data.append({
        'Dağıtım Bölgesi': row[0],
        'Etso Kodu': row[1],
        'Müşteri': row[2],
        'Tarife Grubu': row[3],
        'AG OG': row[4],
        'TERİM': row[5],
        'Güç kW': row[6],
        'KURULU GÜÇ': row[7],
        'Aktif Enerji Tüketim (kWh)': row[8],
        'Dağıtım Bedeli(TL)': row[9],
        'Güç Bedeli(TL)': row[10],
        'Güç Aşım Bedeli (TL)': row[11],
        'Reaktif Bedel (TL)': row[12]
    })

# Filter for positive bedeller only
extracted_filtered = [r for r in extracted_data if r.get('Dağıtım Bedeli(TL)', 0) >= 0]
print(f'Extracted data total: {len(extracted_data)}')
print(f'Extracted (positive only): {len(extracted_filtered)}')

# Load reference
ref_wb = openpyxl.load_workbook('Dağıtımın Kestiği Faturalar Özet.xlsx')
ref_ws = ref_wb['Dağıtımın Kestiği']
ref_rows = list(ref_ws.iter_rows(values_only=True))
reference_data = []
for row in ref_rows[1:]:
    if row[0] and row[0] != 'Dağıtım Bölgesi':
        reference_data.append({
            'Dağıtım Bölgesi': normalize_region_name(row[0], to_format='extraction'),
            'Etso Kodu': row[1],
            'Müşteri': row[2],
            'Tarife Grubu': row[3],
            'AG OG': row[4],
            'TERİM': row[5],
            'Güç kW': row[6],
            'KURULU GÜÇ': row[7],
            'Aktif Enerji Tüketim (kWh)': row[8],
            'Dağıtım Bedeli(TL)': row[9],
            'Güç Bedeli(TL)': row[10],
            'Güç Aşım Bedeli (TL)': row[11],
            'Reaktif Bedel (TL)': row[12]
        })

print(f'Reference data: {len(reference_data)}')

# Check region distribution
print("\n=== REGION DISTRIBUTION ===")
ref_regions = {}
for r in reference_data:
    region = r['Dağıtım Bölgesi']
    ref_regions[region] = ref_regions.get(region, 0) + 1

ext_regions = {}
for r in extracted_filtered:
    region = r['Dağıtım Bölgesi']
    ext_regions[region] = ext_regions.get(region, 0) + 1

print("\nReference by region:")
for region, count in sorted(ref_regions.items()):
    ext_count = ext_regions.get(region, 0)
    print(f"  {region}: Ref={count}, Ext={ext_count}, Diff={count-ext_count}")

print("\nExtra in extraction by region:")
for region, count in sorted(ext_regions.items()):
    ref_count = ref_regions.get(region, 0)
    if count > ref_count:
        print(f"  {region}: Ext={count}, Ref={ref_count}, Extra={count-ref_count}")

# Check for negative bedeller still in comparison
negative_in_filtered = [r for r in extracted_filtered if r.get('Dağıtım Bedeli(TL)', 0) < 0]
print(f"\nNegative bedeller in filtered data (should be 0): {len(negative_in_filtered)}")

# Check region normalization - reference to extraction mapping
print("\n=== REGION MAPPING ISSUES ===")
print("\nRegions in reference but missing in extraction:")
missing_in_ext = set(ref_regions.keys()) - set(ext_regions.keys())
for r in sorted(missing_in_ext):
    print(f"  {r}: {ref_regions[r]} records")

print("\nRegions in extraction but not in reference:")
missing_in_ref = set(ext_regions.keys()) - set(ref_regions.keys())
for r in sorted(missing_in_ref):
    print(f"  {r}: {ext_regions[r]} records")
