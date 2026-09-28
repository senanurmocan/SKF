"""
Core Execution Engine for EDAS Bill Extraction & Excel Output Generation.
"""

import os
import openpyxl
from config.constants import (
    STANDARD_COLUMNS, SPECIAL_HEADER_MAPPING, FIELD_SYNONYMS,
    ABSURD_VALUE_THRESHOLD, COMMA_STYLE_FORMAT, REGION_NAME_MAPPING
)
from core.normalizer import (
    clean_turkish_number, normalize_region_name, normalize_header,
    normalize_etso_kodu
)
from core.mapping import parse_mapping_file, apply_note_rules, extract_file_filter
from core.parser import (
    detect_file_format, parse_excel_file, parse_xml_file, parse_html_file,
    read_xml_content, read_html_content
)
from core.aggregator import aggregate_etso_records


def resolve_column_indices(headers, region_mapping, region_name=None):
    """
    Dynamically resolve header column indices for a given table header array.
    1. Checks region-specific SPECIAL_HEADER_MAPPING override
    2. Matches target mapping header name
    3. Falls back to letter-to-index mapping
    4. Falls back to synonym dictionary (FIELD_SYNONYMS)
    """
    if not headers:
        return {f: region_mapping.get(f'{f}_index', -1) for f in FIELD_SYNONYMS.keys()}

    normalized_headers = [normalize_header(h) for h in headers]
    resolved = {}

    fields = ['etso', 'musteri', 'tarife', 'ag_og', 'terim', 'güç_kw', 'kurulu_güç',
              'aktif_enerji', 'trafo_kaybı', 'dagitim_bedeli', 'güç_bedeli',
              'güç_aşım', 'reaktif', 'reaktif2', 'reaktif_tenzil', 'reaktif_tenzil2']

    norm_reg = normalize_region_name(region_name or region_mapping.get('region_name', ''), to_format='reference')

    for field in fields:
        header_name = region_mapping.get(field)
        fixed_idx = region_mapping.get(f'{field}_index', -1)
        found_idx = -1

        # Check SPECIAL_HEADER_MAPPING override
        if norm_reg in SPECIAL_HEADER_MAPPING and field in SPECIAL_HEADER_MAPPING[norm_reg]:
            sp_val = SPECIAL_HEADER_MAPPING[norm_reg][field]
            if sp_val is None:
                found_idx = -2
            elif isinstance(sp_val, int):
                found_idx = sp_val
            else:
                norm_sp = normalize_header(sp_val)
                if norm_sp in normalized_headers:
                    found_idx = normalized_headers.index(norm_sp)

        if found_idx == -1 and header_name and not pd_isna(header_name) and str(header_name).strip() != '-':
            norm_target = normalize_header(header_name)
            if norm_target and norm_target in normalized_headers:
                found_idx = normalized_headers.index(norm_target)

        if found_idx == -1 and 0 <= fixed_idx < len(headers):
            found_idx = fixed_idx

        if found_idx == -1:
            for syn in FIELD_SYNONYMS.get(field, []):
                if syn in normalized_headers:
                    found_idx = normalized_headers.index(syn)
                    break

        if found_idx == -2:
            found_idx = -1

        resolved[field] = found_idx

    return resolved


def pd_isna(val):
    return val is None or val == '' or str(val).strip().lower() in ('none', 'nan')


def extract_value_from_row(row, column_index, field_name=None, region_name=None):
    """Extract cell value safely with specific handling for ETSO & Yeşilırmak Güç kW strings."""
    if 0 <= column_index < len(row):
        value = row[column_index]
        if field_name == 'etso':
            if isinstance(value, str):
                return value.strip() if value.strip() else None
            if isinstance(value, (int, float)):
                return str(int(value)) if value == int(value) else str(value)
            return value if value else None

        norm_reg = normalize_region_name(region_name or '', to_format='reference') if region_name else ''
        if ('yeşilırmak' in str(region_name).lower() or 'yeşilırmak' in str(norm_reg).lower() or 'yesilirmak' in normalize_header(region_name)) and field_name == 'Güç kW':
            if isinstance(value, str):
                val_str = value.strip()
                if val_str and val_str.lower() not in ('boş', 'none', 'null'):
                    if val_str.endswith('-'):
                        val_str = '-' + val_str[:-1].strip()
                    return val_str

        return clean_turkish_number(value, region_name=region_name, field_name=field_name)
    return 0


def read_excel_content(file_info, region_mapping, region_name=None):
    """Extract data rows from XLS/XLSX file."""
    data_rows = []
    wb = file_info.get('workbook')
    if not wb:
        return data_rows

    is_xlsx = file_info.get('type') == 'xlsx'
    sheet_names = wb.sheetnames if is_xlsx else wb.sheet_names()

    for sheet_name in sheet_names:
        if is_xlsx:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
        else:
            ws = wb.sheet_by_name(sheet_name)
            rows = [ws.row_values(r) for r in range(ws.nrows)]

        if not rows or len(rows) < 2:
            continue

        headers = rows[0]
        header_row_index = 0

        row0_vals = [h for h in headers if h is not None and str(h).strip()]
        if len(row0_vals) <= 2:
            for i, row in enumerate(rows):
                r_vals = [h for h in row if h is not None and str(h).strip()]
                if len(r_vals) > 2:
                    headers = row
                    header_row_index = i
                    break

        if not any(h is not None and str(h).strip() for h in headers):
            continue

        indices = resolve_column_indices(headers, region_mapping, region_name=region_name)
        etso_idx = indices['etso']
        musteri_idx = indices['musteri']
        tarife_idx = indices['tarife']
        ag_og_idx = indices['ag_og']
        term_idx = indices['terim']
        guc_kw_idx = indices['güç_kw']
        kurulu_guc_idx = indices['kurulu_güç']
        aktif_idx = indices['aktif_enerji']
        dagitim_idx = indices['dagitim_bedeli']
        guc_bedeli_idx = indices['güç_bedeli']
        guc_asim_idx = indices['güç_aşım']
        trafo_kaybi_idx = indices['trafo_kaybı']

        pr_sum_rule = region_mapping.get('rules', {}).get('active_energy_rule') == 'P + R sum'
        is_single_reactive = region_mapping.get('rules', {}).get('reactive_total_rule') == 'single_column'
        norm_rn = normalize_region_name(region_name or '', to_format='reference')
        no_tenzil_regions = ['Aras EDAŞ', 'Çoruh EDAŞ', 'Dicle Edaş', 'Fırat Edaş', 'Vangölü Edaş']

        if is_single_reactive or norm_rn in no_tenzil_regions:
            reaktif_fields = ['reaktif', 'reaktif2']
        else:
            reaktif_fields = ['reaktif', 'reaktif2', 'reaktif_tenzil', 'reaktif_tenzil2']

        if norm_rn in ['Trakya EDAŞ', 'Trakya Edaş', 'TREDAŞ']:
            ilk_reaktif_fields = ['reaktif_tenzil2']
        else:
            ilk_reaktif_fields = ['reaktif_tenzil', 'reaktif_tenzil2']

        for row in rows[header_row_index + 1:]:
            if not row or all(v is None for v in row):
                continue

            reaktif_toplam = 0.0
            sakarya_x = False
            for f_key in reaktif_fields:
                idx = indices.get(f_key, -1)
                if 0 <= idx < len(row):
                    val = row[idx]
                    if val is not None:
                        if str(val).strip().upper() == 'X':
                            sakarya_x = True
                        else:
                            reaktif_toplam += clean_turkish_number(val, region_name=region_name, field_name='reaktif')

            ilk_reaktif_toplam = 0.0
            for f_key in ilk_reaktif_fields:
                idx = indices.get(f_key, -1)
                if 0 <= idx < len(row):
                    val = row[idx]
                    if val is not None:
                        if str(val).strip().upper() == 'X':
                            sakarya_x = True
                        else:
                            ilk_reaktif_toplam += clean_turkish_number(val, region_name=region_name, field_name='reaktif')

            guc_val = extract_value_from_row(row, guc_kw_idx, field_name='Güç kW', region_name=region_name)
            kurulu_val = extract_value_from_row(row, kurulu_guc_idx, field_name='kurulu_güç', region_name=region_name)
            aktif_val = extract_value_from_row(row, aktif_idx, field_name='aktif_enerji', region_name=region_name)
            trafo_val = extract_value_from_row(row, trafo_kaybi_idx, field_name='trafo_kaybı', region_name=region_name)
            dagitim_val = extract_value_from_row(row, dagitim_idx, field_name='dagitim_bedeli', region_name=region_name)
            guc_bedeli_val = extract_value_from_row(row, guc_bedeli_idx, field_name='güç_bedeli', region_name=region_name)
            guc_asim_val = extract_value_from_row(row, guc_asim_idx, field_name='güç_aşım', region_name=region_name)

            data_row = {
                'Dağıtım Bölgesi': region_mapping.get('region_name', ''),
                'Etso Kodu': extract_value_from_row(row, etso_idx, field_name='etso', region_name=region_name),
                'Müşteri': row[musteri_idx] if 0 <= musteri_idx < len(row) else None,
                'Tarife Grubu': row[tarife_idx] if 0 <= tarife_idx < len(row) else None,
                'AG OG': row[ag_og_idx] if 0 <= ag_og_idx < len(row) else None,
                'TERİM': row[term_idx] if 0 <= term_idx < len(row) else None,
                'Güç kW': guc_val,
                'KURULU GÜÇ': kurulu_val,
                'Aktif Enerji Tüketim (kWh)': (aktif_val + trafo_val) if pr_sum_rule and trafo_kaybi_idx >= 0 else aktif_val,
                'Dağıtım Bedeli(TL)': dagitim_val,
                'Güç Bedeli(TL)': guc_bedeli_val,
                'Güç Aşım Bedeli (TL)': guc_asim_val,
                'Reaktif Bedel (TL)': reaktif_toplam,
                'İlk Reaktif': 'X' if sakarya_x else ilk_reaktif_toplam
            }
            data_rows.append(data_row)

    return data_rows


def filter_extracted_data(data, region_name=None):
    """Filter zero/absurd bedel values per region filtering rules."""
    filtered_data = []
    zero_bedel_count = 0
    negative_bedel_count = 0
    akdeniz_filtered_count = 0
    absurd_bedel_count = 0

    norm_region = normalize_region_name(region_name, to_format='reference') if region_name else None

    for row in data:
        dagitim_bedeli = float(row.get('Dağıtım Bedeli(TL)', 0) or 0)
        musteri = row.get('Müşteri')
        is_musteri_empty = (musteri is None or str(musteri).strip() == '' or str(musteri).strip().lower() == 'none')

        if abs(dagitim_bedeli) < 1e-4 and is_musteri_empty:
            zero_bedel_count += 1
            continue

        if dagitim_bedeli < 0:
            negative_bedel_count += 1
            if norm_region == 'Akdeniz EDAŞ':
                akdeniz_filtered_count += 1
                continue

        if abs(dagitim_bedeli) > ABSURD_VALUE_THRESHOLD:
            absurd_bedel_count += 1
            continue

        filtered_data.append(row)

    return filtered_data, zero_bedel_count, negative_bedel_count, akdeniz_filtered_count, absurd_bedel_count


def process_all_regions(base_path, mapping_file_path, status_callback=None):
    """
    Main Orchestration Engine: Processes all 21 EDAŞ folders using SKF mapping file.
    Returns: (all_extracted_data, stats_dict)
    """
    if not os.path.exists(mapping_file_path):
        raise FileNotFoundError(f"Mapping file missing: {mapping_file_path}")

    mapping_data = parse_mapping_file(mapping_file_path)
    all_extracted_data = []
    stats = {
        'total_extracted': 0,
        'zero_filtered': 0,
        'negative_filtered': 0,
        'akdeniz_filtered': 0,
        'absurd_filtered': 0,
        'merged_duplicates': 0,
        'format_adaptations': []
    }

    if status_callback:
        status_callback("Mapping okundu, 21 EDAŞ bölgesi işleniyor...", 5)

    # Dynamically discover region directories inside base_path (excluding system/module folders)
    ignored_folders = {'core', 'config', 'analytics', '__pycache__', '.git', '.system_generated', '.agents'}
    all_subdirs = [
        d for d in os.listdir(base_path)
        if os.path.isdir(os.path.join(base_path, d))
        and d not in ignored_folders
        and not d.startswith('.')
    ]

    total_regions = len(all_subdirs)
    for idx_reg, region_folder in enumerate(all_subdirs):
        region_dir = os.path.join(base_path, region_folder)
        pct = 10 + int(((idx_reg + 1) / max(1, total_regions)) * 80)

        if not os.path.exists(region_dir):
            if status_callback:
                status_callback(f"Klasör bulunamadı atlanıyor: {region_folder}", pct)
            continue

        if status_callback:
            status_callback(f"İşleniyor: {region_folder}", pct)

        region_name = REGION_NAME_MAPPING.get(region_folder, region_folder)
        region_mapping = mapping_data['regions'].get(region_name, mapping_data['regions'].get(region_folder, {}))
        region_mapping['rules'] = apply_note_rules(mapping_data, region_name)
        region_mapping['region_name'] = region_name

        # Requirement 2: Recursive subfolder file discovery via os.walk
        excel_files = []
        for root_dir, dirs, files in os.walk(region_dir):
            for fname in files:
                if fname.endswith(('.xlsx', '.XLSX', '.xls', '.XLS', '.xml', '.html')):
                    excel_files.append(os.path.join(root_dir, fname))

        if not excel_files:
            continue

        region_raw_data = []

        for fpath in excel_files:
            fname = os.path.basename(fpath)
            fformat = detect_file_format(fpath)

            if fformat in ('excel', 'xlsx', 'xls'):
                finfo = parse_excel_file(fpath)
                if finfo and finfo['type'] in ('xlsx', 'xls'):
                    data = read_excel_content(finfo, region_mapping, region_name)
                    region_raw_data.extend(data)
                elif finfo is None:
                    # XML/HTML fallback
                    fformat = detect_file_format(fpath)

            if fformat == 'xml':
                finfo = parse_xml_file(fpath)
                if finfo and finfo['type'] == 'xml':
                    data = read_xml_content(finfo, region_mapping, region_name, resolve_column_indices, extract_value_from_row)
                    region_raw_data.extend(data)

            elif fformat == 'html':
                finfo = parse_html_file(fpath)
                if finfo and finfo['type'] == 'html':
                    data = read_html_content(finfo, region_mapping, region_name, resolve_column_indices, extract_value_from_row)
                    region_raw_data.extend(data)

        filtered_data, zf, nf, af, abf = filter_extracted_data(region_raw_data, region_name=region_name)
        stats['total_extracted'] += len(region_raw_data)
        stats['zero_filtered'] += zf
        stats['negative_filtered'] += nf
        stats['akdeniz_filtered'] += af
        stats['absurd_filtered'] += abf

        all_extracted_data.extend(filtered_data)

    # ETSO Aggregation across all records
    if status_callback:
        status_callback("ETSO kayıtları birleştiriliyor...", 92)

    final_aggregated_data, merge_count = aggregate_etso_records(all_extracted_data)
    stats['merged_duplicates'] = merge_count

    # 21 Region Tracking & Missing Region Audit (Requirement 1: Eşsiz Bölge Sayacı Fix)
    all_canonical_regions = [
        'ADM EDAŞ', 'Akdeniz EDAŞ', 'AKEDAŞ', 'Aras EDAŞ', 'AYEDAŞ',
        'Başkent EDAŞ', 'Boğaziçi EDAŞ', 'Çamlıbel EDAŞ', 'Çoruh EDAŞ',
        'Dicle EDAŞ', 'Fırat EDAŞ', 'Gediz EDAŞ', 'Kayseri ve Civarı',
        'Meram EDAŞ', 'Osmangazi EDAŞ', 'Sakarya EDAŞ', 'Toroslar EDAŞ',
        'Trakya EDAŞ', 'Uludağ EDAŞ', 'Vangölü EDAŞ', 'Yeşilırmak EDAŞ'
    ]

    processed_region_names = set(
        row.get('Dağıtım Bölgesi') for row in final_aggregated_data if row.get('Dağıtım Bölgesi')
    )
    
    successful_regions = []
    missing_regions = []

    for reg in all_canonical_regions:
        norm_reg = normalize_region_name(reg, to_format='reference')
        found = any(normalize_region_name(pr, to_format='reference') == norm_reg for pr in processed_region_names)
        if found:
            successful_regions.append(reg)
        else:
            missing_regions.append(reg)

    stats['successful_regions'] = successful_regions
    stats['missing_regions'] = missing_regions
    stats['read_region_count'] = len(successful_regions)

    if status_callback:
        status_callback("İşlem tamamlandı!", 100)

    return final_aggregated_data, stats


def save_to_excel(data, filepath):
    """Save extracted dataset to Excel with standard headers, KDV & Toplam formulas, Comma Style formatting, and Auto-Fit column widths."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Çıkarılan Veriler'

    ws.append(STANDARD_COLUMNS)
    for col_idx in range(1, len(STANDARD_COLUMNS) + 1):
        ws.cell(1, col_idx).number_format = 'General'

    for idx, row in enumerate(data, start=2):
        ws.append([
            row.get('Dağıtım Bölgesi', ''),
            row.get('Etso Kodu', ''),
            row.get('Müşteri', ''),
            row.get('Tarife Grubu', ''),
            row.get('AG OG', ''),
            row.get('TERİM', ''),
            row.get('Güç kW', ''),
            row.get('KURULU GÜÇ', ''),
            row.get('Aktif Enerji Tüketim (kWh)', ''),
            row.get('Dağıtım Bedeli(TL)', ''),
            row.get('Güç Bedeli(TL)', ''),
            row.get('Güç Aşım Bedeli (TL)', ''),
            row.get('Reaktif Bedel (TL)', ''),
            row.get('İlk Reaktif', ''),
            f"=(J{idx}+K{idx}+L{idx}+M{idx})*0.2",
            f"=SUM(J{idx}:M{idx})+O{idx}",
            None
        ])

        # Formatting
        for c in [1, 3, 4, 5, 6, 17]:
            ws.cell(idx, c).number_format = 'General'

        ws.cell(idx, 2).number_format = '0'

        for c in range(7, 17):
            ws.cell(idx, c).number_format = COMMA_STYLE_FORMAT

    # Sütunları otomatik genişletme
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[col_letter].width = adjusted_width

    wb.save(filepath)
    return filepath
