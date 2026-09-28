"""
File format detection and parsers for Excel (XLS, XLSX), XML (SpreadsheetML), and HTML formats.
"""

import os
import openpyxl
import xlrd
import xml.etree.ElementTree as ET
from html.parser import HTMLParser
from core.normalizer import clean_turkish_number, normalize_region_name


class ExcelHTMLParser(HTMLParser):
    """Custom HTML Parser for Excel HTML format tables."""
    def __init__(self):
        super().__init__()
        self.tables = []
        self.current_table = None
        self.current_row = None
        self.in_cell = False
        self.cell_data = []

    def handle_starttag(self, tag, attrs):
        if tag == 'table':
            self.current_table = []
            self.current_row = None
        elif tag == 'tr':
            self.current_row = []
        elif tag in ('td', 'th'):
            self.in_cell = True
            self.cell_data = []

    def handle_endtag(self, tag):
        if tag == 'table':
            if self.current_table:
                self.tables.append(self.current_table)
            self.current_table = None
            self.current_row = None
        elif tag == 'tr':
            if self.current_row is not None and self.current_table is not None:
                self.current_table.append(self.current_row)
                self.current_row = None
        elif tag in ('td', 'th'):
            self.in_cell = False
            cell_text = ''.join(self.cell_data).strip()
            if self.current_row is not None:
                self.current_row.append(cell_text if cell_text else None)

    def handle_data(self, data):
        if self.in_cell:
            self.cell_data.append(data)


def detect_file_format(filepath):
    """Detect file format based on content signature (magic bytes)."""
    try:
        with open(filepath, 'rb') as f:
            header = f.read(20)

        if header.startswith(b'<?xml'):
            return 'xml'
        elif b'<html xm' in header.lower() or header.startswith(b'<html'):
            return 'html'
        elif header.startswith(b'\x50\x4B\x03\x04'):
            return 'xlsx'  # ZIP format
        elif header.startswith(b'\xD0\xCF\x11\xE0'):
            return 'xls'   # OLE2 format
        else:
            return 'excel'
    except Exception as e:
        return 'unknown'


def parse_excel_file(filepath):
    """Parse standard Excel file (XLS / XLSX) with fallback format checks."""
    try:
        with open(filepath, 'rb') as f:
            header = f.read(20)

        if header.startswith(b'<?xml'):
            return None  # Should use parse_xml_file
        elif b'<html xm' in header.lower() or header.startswith(b'<html'):
            return None  # Should use parse_html_file
        elif filepath.endswith(('.xlsx', '.XLSX')):
            wb = openpyxl.load_workbook(filepath, data_only=True)
            return {'type': 'xlsx', 'workbook': wb}
        elif filepath.endswith(('.xls', '.XLS')):
            try:
                wb = xlrd.open_workbook(filepath)
                return {'type': 'xls', 'workbook': wb}
            except Exception as xlrd_e:
                with open(filepath, 'rb') as f:
                    content = f.read(500)
                if b'<html xm' in content.lower() or content.startswith(b'<html'):
                    return None
                elif content.startswith(b'<?xml') or b'<html' in content:
                    return None
                raise xlrd_e
    except Exception as e:
        return {'type': 'error', 'error': str(e)}
    return None


def parse_xml_file(filepath):
    """Parse Excel 2003 XML SpreadsheetML format."""
    try:
        tree = ET.parse(filepath)
        root = tree.getroot()
        return {'type': 'xml', 'root': root}
    except Exception as e:
        return {'type': 'error', 'error': str(e)}


def parse_html_file(filepath):
    """Parse Excel HTML format."""
    try:
        with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
            content = f.read()
        return {'type': 'html', 'content': content}
    except Exception as e:
        return {'type': 'error', 'error': str(e)}


def read_xml_content(file_info, region_mapping, region_name, resolve_column_indices_func, extract_value_func):
    """Extract rows from XML SpreadsheetML file."""
    data_rows = []
    if file_info.get('type') != 'xml':
        return data_rows

    root = file_info['root']
    ns = {'ss': 'urn:schemas-microsoft-com:office:spreadsheet'}
    worksheets = root.findall('.//ss:Worksheet', ns)

    for worksheet in worksheets:
        table = worksheet.find('ss:Table', ns)
        if table is None:
            continue

        rows = list(table.findall('ss:Row', ns))
        if not rows:
            continue

        headers = None
        header_row_index = 0

        # Read first row
        headers_row = rows[0]
        headers = []
        for cell in headers_row.findall('ss:Cell', ns):
            data_elem = cell.find('ss:Data', ns)
            headers.append(data_elem.text if data_elem is not None else None)

        row0_vals = [h for h in headers if h is not None and str(h).strip()]
        if len(row0_vals) <= 2:
            for i, row in enumerate(rows):
                row_headers = []
                for cell in row.findall('ss:Cell', ns):
                    data_elem = cell.find('ss:Data', ns)
                    row_headers.append(data_elem.text if (data_elem is not None and data_elem.text) else None)
                if len([h for h in row_headers if h is not None and str(h).strip()]) > 2:
                    headers = row_headers
                    header_row_index = i
                    break

        if not headers or not any(h is not None and str(h).strip() for h in headers):
            continue

        indices = resolve_column_indices_func(headers, region_mapping, region_name=region_name)
        _process_rows(rows[header_row_index + 1:], indices, ns, True, region_mapping, region_name, data_rows, extract_value_func)

    return data_rows


def read_html_content(file_info, region_mapping, region_name, resolve_column_indices_func, extract_value_func):
    """Extract rows from Excel HTML file."""
    data_rows = []
    if file_info.get('type') != 'html':
        return data_rows

    parser = ExcelHTMLParser()
    parser.feed(file_info['content'])

    if not parser.tables:
        return data_rows

    table = max(parser.tables, key=len)
    if not table or len(table) < 2:
        return data_rows

    headers = table[0]
    header_row_index = 0

    row0_vals = [h for h in headers if h is not None and str(h).strip()]
    if len(row0_vals) <= 2:
        for i, row in enumerate(table):
            r_vals = [h for h in row if h is not None and str(h).strip()]
            if len(r_vals) > 2:
                headers = row
                header_row_index = i
                break

    if not any(h is not None and str(h).strip() for h in headers):
        return data_rows

    indices = resolve_column_indices_func(headers, region_mapping, region_name=region_name)
    _process_rows(table[header_row_index + 1:], indices, None, False, region_mapping, region_name, data_rows, extract_value_func)

    return data_rows


def _process_rows(raw_rows, indices, ns, is_xml, region_mapping, region_name, data_rows, extract_value_func):
    """Helper row extractor for XML and HTML tables."""
    etso_idx = indices['etso']
    musteri_idx = indices['musteri']
    tarife_idx = indices['tarife']
    ag_og_idx = indices['ag_og']
    term_idx = indices['terim']
    guc_kw_idx = indices['güç_kw']
    kurulu_guc_idx = indices['kurulu_güç']
    aktif_idx = indices['aktif_enerji']
    dagitim_idx = indices['dagitim_bedeli']
    guc_bedeli_idx = indices['güç_bedeli']
    guc_asim_idx = indices['güç_aşım']
    trafo_kaybi_idx = indices['trafo_kaybı']

    pr_sum_rule = region_mapping.get('rules', {}).get('active_energy_rule') == 'P + R sum'
    is_single_reactive = region_mapping.get('rules', {}).get('reactive_total_rule') == 'single_column'
    norm_rn = normalize_region_name(region_name or '', to_format='reference')
    no_tenzil_regions = ['Aras EDAŞ', 'Çoruh EDAŞ', 'Dicle Edaş', 'Fırat Edaş', 'Vangölü Edaş']

    if is_single_reactive or norm_rn in no_tenzil_regions:
        reaktif_fields = ['reaktif', 'reaktif2']
    else:
        reaktif_fields = ['reaktif', 'reaktif2', 'reaktif_tenzil', 'reaktif_tenzil2']

    if norm_rn in ['Trakya EDAŞ', 'Trakya Edaş', 'TREDAŞ']:
        ilk_reaktif_fields = ['reaktif_tenzil2']
    else:
        ilk_reaktif_fields = ['reaktif_tenzil', 'reaktif_tenzil2']

    for row_elem in raw_rows:
        if is_xml:
            cells = row_elem.findall('ss:Cell', ns)
            row = [c.find('ss:Data', ns).text if c.find('ss:Data', ns) is not None else None for c in cells]
        else:
            row = row_elem

        if not row or all(v is None for v in row):
            continue

        reaktif_toplam = 0.0
        sakarya_x = False
        for f_key in reaktif_fields:
            idx = indices.get(f_key, -1)
            if 0 <= idx < len(row):
                val = row[idx]
                if val is not None:
                    if str(val).strip().upper() == 'X':
                        sakarya_x = True
                    else:
                        reaktif_toplam += clean_turkish_number(val, region_name=region_name, field_name='reaktif')

        ilk_reaktif_toplam = 0.0
        for f_key in ilk_reaktif_fields:
            idx = indices.get(f_key, -1)
            if 0 <= idx < len(row):
                val = row[idx]
                if val is not None:
                    if str(val).strip().upper() == 'X':
                        sakarya_x = True
                    else:
                        ilk_reaktif_toplam += clean_turkish_number(val, region_name=region_name, field_name='reaktif')

        guc_val = extract_value_func(row, guc_kw_idx, field_name='Güç kW', region_name=region_name)
        kurulu_val = extract_value_func(row, kurulu_guc_idx, field_name='kurulu_güç', region_name=region_name)
        aktif_val = extract_value_func(row, aktif_idx, field_name='aktif_enerji', region_name=region_name)
        trafo_val = extract_value_func(row, trafo_kaybi_idx, field_name='trafo_kaybı', region_name=region_name)
        dagitim_val = extract_value_func(row, dagitim_idx, field_name='dagitim_bedeli', region_name=region_name)
        guc_bedeli_val = extract_value_func(row, guc_bedeli_idx, field_name='güç_bedeli', region_name=region_name)
        guc_asim_val = extract_value_func(row, guc_asim_idx, field_name='güç_aşım', region_name=region_name)

        data_row = {
            'Dağıtım Bölgesi': region_mapping.get('region_name', ''),
            'Etso Kodu': extract_value_func(row, etso_idx, field_name='etso', region_name=region_name),
            'Müşteri': row[musteri_idx] if 0 <= musteri_idx < len(row) else None,
            'Tarife Grubu': row[tarife_idx] if 0 <= tarife_idx < len(row) else None,
            'AG OG': row[ag_og_idx] if 0 <= ag_og_idx < len(row) else None,
            'TERİM': row[term_idx] if 0 <= term_idx < len(row) else None,
            'Güç kW': guc_val,
            'KURULU GÜÇ': kurulu_val,
            'Aktif Enerji Tüketim (kWh)': (aktif_val + trafo_val) if pr_sum_rule and trafo_kaybi_idx >= 0 else aktif_val,
            'Dağıtım Bedeli(TL)': dagitim_val,
            'Güç Bedeli(TL)': guc_bedeli_val,
            'Güç Aşım Bedeli (TL)': guc_asim_val,
            'Reaktif Bedel (TL)': reaktif_toplam,
            'İlk Reaktif': 'X' if sakarya_x else ilk_reaktif_toplam
        }
        data_rows.append(data_row)
