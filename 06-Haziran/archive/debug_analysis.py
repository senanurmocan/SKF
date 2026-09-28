#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Debug analysis for discrepancy report
"""
from pathlib import Path

import openpyxl
import sys
sys.path.append(str(Path(__file__).resolve().parents[1]))
from new_mapping_parser import parse_mapping_file, apply_note_rules, excel_column_letter_to_index

BASE_PATH = Path(__file__).resolve().parents[1]
REFERENCE_FILE = BASE_PATH / 'Dağıtımın Kestiği Faturalar Özet.xlsx'
EXTRACTED_FILE = BASE_PATH / 'Çıkarılan_Veriler.xlsx'

def normalize_region_name(region_name, to_format='reference'):
    """Region ismini standartlaştırmak için fonksiyon"""
    if region_name is None:
        return None

    region_name = str(region_name).strip().upper()

    REGION_NORMALIZATION_MAPPING = {
        "AKEDAŞ": "AKEDAŞ ( Göksu EDAŞ )",
        "ADM EDAŞ": "Aydem EDAŞ",
        "AYEDAŞ": "Ayedaş",
        "DİCLE EDAŞ": "Dicle Edaş",
        "FIRAT EDAŞ": "Fırat Edaş",
        "KAYSERI VE CIVARI": "Kayseri EDAŞ",
        "KCETAŞ": "Kayseri EDAŞ",
        "VANGÖLÜ EDAŞ": "Vangölü Edaş",
        "BOĞAZİÇİ EDAŞ": "Boğaziçi EDAŞ",
        "BAŞKENT EDAŞ": "Başkent EDAŞ",
        "SEDAŞ": "Sakarya EDAŞ",
        "TREDAŞ": "Trakya EDAŞ",
        "UEDAŞ": "Uludağ EDAŞ",
        "ÇAMLIBEL EDAŞ": "Çamlıbel EDAŞ",
        "YEŞİLIRMAK EDAŞ": "Yeşilırmak EDAŞ",
        "GDZ EDAŞ": "Gediz EDAŞ",
        "MERAM": "Meram EDAŞ",
        "OEDAŞ": "Osmangazi EDAŞ",
        "TOROSLAR EDAŞ": "Toroslar EDAŞ",
        "ÇORUH EDAŞ": "Çoruh EDAŞ",
        "ARAS EDAŞ": "Aras EDAŞ",
        "Gediz EDAŞ": "Gediz EDAŞ",
        "Meram EDAŞ": "Meram EDAŞ",
        "Osmangazi EDAŞ": "Osmangazi EDAŞ",
        "Sakarya EDAŞ": "Sakarya EDAŞ",
        "Trakya EDAŞ": "Trakya EDAŞ",
        "Uludağ EDAŞ": "Uludağ EDAŞ",
        "Çamlıbel EDAŞ": "Çamlıbel EDAŞ",
        "Vangölü EDAŞ": "Vangölü Edaş",
        "Yeşilırmak EDAŞ": "Yeşilırmak EDAŞ",
        "Başkent EDAŞ": "Başkent EDAŞ",
        "Toroslar EDAŞ": "Toroslar EDAŞ",
        "Fırat EDAŞ": "Fırat Edaş",
        "Dicle EDAŞ": "Dicle Edaş",
        "Aras EDAŞ": "Aras EDAŞ",
        "Çoruh EDAŞ": "Çoruh EDAŞ",
    }

    return REGION_NORMALIZATION_MAPPING.get(region_name, region_name)

# Check region values
print("=== REFERENCE FILE REGION VALUES ===")
wb = openpyxl.load_workbook(REFERENCE_FILE, data_only=True)
ws = wb.active
regions_ref = {}
for row in list(ws.iter_rows(values_only=True))[1:]:
    if row and row[0]:
        r = str(row[0]).strip()
        regions_ref[r] = regions_ref.get(r, 0) + 1
for r, count in sorted(regions_ref.items(), key=lambda x: x[1], reverse=True):
    print(f"  {r}: {count}")

print("\n=== EXTRACTED FILE REGION VALUES ===")
wb2 = openpyxl.load_workbook(EXTRACTED_FILE, data_only=True)
ws2 = wb2.active
regions_ext = {}
for row in list(ws2.iter_rows(values_only=True))[1:]:
    if row and row[0]:
        r = str(row[0]).strip()
        regions_ext[r] = regions_ext.get(r, 0) + 1
for r, count in sorted(regions_ext.items(), key=lambda x: x[1], reverse=True):
    print(f"  {r}: {count}")

print("\n=== REGION MAPPING TEST ===")
test_regions = ["AYDEM EDAŞ", "ADM EDAŞ", "Aydem EDAŞ"]
for tr in test_regions:
    print(f"  '{tr}' -> '{normalize_region_name(tr)}'")

print("\n=== Etso Kodu comparison test ===")
# Get first few Etso values
wb = openpyxl.load_workbook(REFERENCE_FILE, data_only=True)
ws = wb.active
ref_etso = []
for row in list(ws.iter_rows(values_only=True))[1:6]:
    if row and row[1]:
        ref_etso.append((type(row[1]).__name__, row[1]))

wb2 = openpyxl.load_workbook(EXTRACTED_FILE, data_only=True)
ws2 = wb2.active
ext_etso = []
for row in list(ws2.iter_rows(values_only=True))[1:6]:
    if row and row[1]:
        ext_etso.append((type(row[1]).__name__, row[1]))

print("Reference Etso values (type, value):")
for t, v in ref_etso:
    print(f"  {t}: {v}")

print("Extracted Etso values (type, value):")
for t, v in ext_etso:
    print(f"  {t}: {v}")

print("\n=== Sample data comparison ===")
print("First row from reference:")
wb = openpyxl.load_workbook(REFERENCE_FILE, data_only=True)
ws = wb.active
row = list(ws.iter_rows(values_only=True))[1]
print(f"  Region: {row[0]} (type: {type(row[0])})")
print(f"  Etso: {row[1]} (type: {type(row[1])})")
print(f"  Customer: {row[2]} (type: {type(row[2])})")
print(f"  Bedel: {row[9]} (type: {type(row[9])})")

print("\nFirst row from extracted:")
wb2 = openpyxl.load_workbook(EXTRACTED_FILE, data_only=True)
ws2 = wb2.active
row2 = list(ws2.iter_rows(values_only=True))[1]
print(f"  Region: {row2[0]} (type: {type(row2[0])})")
print(f"  Etso: {row2[1]} (type: {type(row2[1])})")
print(f"  Customer: {row2[2]} (type: {type(row2[2])})")
print(f"  Bedel: {row2[9]} (type: {type(row2[9])})")
