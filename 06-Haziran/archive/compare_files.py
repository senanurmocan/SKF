from pathlib import Path
import openpyxl

# Load both files
original_file = Path(__file__).resolve().parents[1] / 'SKF Başlıkları.xlsx'
current_file = Path(__file__).resolve().parents[1] / 'SKF Başlıkları Güncel.xlsx'

wb_orig = openpyxl.load_workbook(original_file)
ws_orig = wb_orig["Başlıklar"]

wb_curr = openpyxl.load_workbook(current_file)
ws_curr = wb_curr["Başlıklar"]

print("=== Regions in Original (SKF Başlıkları.xlsx) ===")
orig_regions = []
for row_idx in range(1, ws_orig.max_row + 1):
    region_name = ws_orig.cell(row=row_idx, column=1).value
    if region_name:
        orig_regions.append((row_idx, region_name))
        print(f"Row {row_idx}: {region_name}")

print("\n=== Regions in Current (SKF Başlıkları Güncel.xlsx) ===")
curr_regions = []
for row_idx in range(1, ws_curr.max_row + 1):
    region_name = ws_curr.cell(row=row_idx, column=1).value
    if region_name:
        curr_regions.append((row_idx, region_name))
        print(f"Row {row_idx}: {region_name}")

# Compare
print("\n=== Comparison ===")
orig_names = [r[1] for r in orig_regions]
curr_names = [r[1] for r in curr_regions]

print(f"Original has {len(orig_names)} regions")
print(f"Current has {len(curr_names)} regions")

# Find differences
orig_set = set(orig_names)
curr_set = set(curr_names)

missing_in_current = orig_set - curr_set
extra_in_current = curr_set - orig_set

if missing_in_current:
    print(f"\nMissing in current: {missing_in_current}")
if extra_in_current:
    print(f"\nExtra in current: {extra_in_current}")

# Check column structure
print("\n=== Column Structure Check ===")
print("Original file column 1-5:")
for row_idx in range(1, min(4, ws_orig.max_row + 1)):
    cols = []
    for col_idx in range(1, 6):
        cols.append(str(ws_orig.cell(row=row_idx, column=col_idx).value))
    print(f"Row {row_idx}: {' | '.join(cols)}")

print("\nCurrent file column 1-5:")
for row_idx in range(1, min(4, ws_curr.max_row + 1)):
    cols = []
    for col_idx in range(1, 6):
        cols.append(str(ws_curr.cell(row=row_idx, column=col_idx).value))
    print(f"Row {row_idx}: {' | '.join(cols)}")
