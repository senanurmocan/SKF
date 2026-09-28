#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
SKF Başlıkları Güncel.xlsx dosyasına Yeşilırmak EDAŞ ekler
"""
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

# Read current file
file_path = Path(__file__).resolve().parents[1] / 'SKF Başlıkları Güncel.xlsx'

# Load workbook to manipulate sheets
wb = load_workbook(file_path)

# Get Başlıklar sheet
ws = wb['Başlıklar']

# Yeşilırmak EDAŞ row - based on the alternating structure
# Row should have: Region name, Etso header, column letter, header, column letter, etc.
# Using the same pattern as other regions (based on Vangölü row which is row 20)

# Read existing row 20 (Vangölü) to copy structure
row_20_data = []
for cell in ws[20]:
    row_20_data.append(cell.value)

print("Row 20 (Vangölü) structure:")
print(row_20_data[:20])

# Yeşilırmak row data - using same pattern as Vangölü
# The region name is at index 0, Etso header at index 1
yesilirmak_row = row_20_data.copy()
yesilirmak_row[0] = 'Yeşilırmak EDAŞ'

print("\nYeşilırmak row to add:")
print(yesilirmak_row[:20])

# Insert as new row 21
ws.append(yesilirmak_row)

# Save
wb.save(file_path)
print(f"\n✓ Yeşilırmak EDAŞ added as row 21")
print(f"✓ Total rows: {ws.max_row}")

# Verify
wb2 = load_workbook(file_path)
ws2 = wb2['Başlıklar']
print("\nVerified regions:")
for i in range(1, ws2.max_row + 1):
    val = ws2.cell(row=i, column=1).value
    if val:
        print(f"  Row {i}: {val}")

wb2.close()
wb.close()
