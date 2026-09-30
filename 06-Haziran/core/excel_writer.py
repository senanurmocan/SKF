# -*- coding: utf-8 -*-
"""
Excel Writer Module for EDAS Bill Extraction & Analysis System.

This module provides functions to format and save extracted bill data to Excel
workbooks with standardized styling, accounting number formats, dynamic formulas
for KDV and totals, and auto-adjusted column dimensions.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from config.constants import STANDARD_COLUMNS, EXTENDED_OUTPUT_COLUMNS

try:
    from core.number_cleaner import _format_guc_kw
except ImportError:
    try:
        from number_cleaner import _format_guc_kw
    except ImportError:
        def _format_guc_kw(value):
            """Yeşilırmak EDAŞ Güç kW string değerlerini (ör. '500,00000000000000') 2 ondalıklı float'a dönüştür."""
            if value is None or value == '':
                return ''
            if isinstance(value, (int, float)):
                return round(float(value), 2)
            val_str = str(value).strip()
            if not val_str or val_str.lower() in ('none', 'boş'):
                return ''
            # Türkçe virgüllü string veya düz sayı
            try:
                cleaned = val_str.replace(' ', '').replace('.', '').replace(',', '.')
                return round(float(cleaned), 2)
            except (ValueError, TypeError):
                return val_str  # Dönüştürülemezse orijinali bırak


def save_to_excel(data, filepath):
    """Veriyi Excel'e kaydet - Formüller, Sayı Biçimleri ve Başlık Biçimlendirmesi ile"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Çıkarılan Veriler'
    
    extended_output_enabled = any(
        any(field in row for field in EXTENDED_OUTPUT_COLUMNS)
        for row in data
    )
    output_columns = list(STANDARD_COLUMNS)
    if extended_output_enabled:
        output_columns.extend(EXTENDED_OUTPUT_COLUMNS)

    # Header satırı ekle
    ws.append(output_columns)

    # Başlık satırı biçimlendirmesi: Kalın, CK Mavi (#305496) arka plan, beyaz yazı, ortalama
    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(fill_type='solid', fgColor='305496')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=False)
    data_align_left = Alignment(horizontal='left', vertical='center')
    for col_idx in range(1, len(output_columns) + 1):
        cell = ws.cell(1, col_idx)
        cell.number_format = 'General'
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    # Başlığı dondur (A2'den itibaren kaydır)
    ws.freeze_panes = 'A2'

    # Otomatik filtre (tüm sütunlara)
    last_col_letter = ws.cell(1, len(output_columns)).column_letter
    ws.auto_filter.ref = f'A1:{last_col_letter}1'

    COMMA_STYLE_FORMAT = '_-* #,##0.00_-;-* #,##0.00_-;_-* "-"??_-;_-@_-'
    
    # Veri satırları
    for idx, row in enumerate(data, start=2):
        output_row = [
            row.get('Dağıtım Bölgesi', ''),
            row.get('Etso Kodu', ''),
            row.get('Müşteri', ''),
            row.get('Tarife Grubu', ''),
            row.get('AG OG', ''),
            row.get('TERİM', ''),
            _format_guc_kw(row.get('Güç kW', '')),
            row.get('KURULU GÜÇ', ''),
            row.get('Aktif Enerji Tüketim (kWh)', ''),
            row.get('Dağıtım Bedeli(TL)', ''),
            row.get('Güç Bedeli(TL)', ''),
            row.get('Güç Aşım Bedeli (TL)', ''),
            row.get('Reaktif Bedel (TL)', ''),
            f"=J{idx}+K{idx}+L{idx}+M{idx}",           # N: KDV Matrahı (TL)
            f"=N{idx}*0.2",                              # O: KDV
            f"=N{idx}+O{idx}+Q{idx}",                    # P: Toplam (TL)
            row.get('İlk Reaktif', ''),                  # Q: İlk Reaktif Bedeli (TL)
            row.get("Sayax'a Atılacak Tarife", '')       # R: Sayax'a Atılacak Tarife
        ]
        if extended_output_enabled:
            output_row.extend([
                row.get('Standart Dışı Tutar (TL)', 0),
                row.get('Tazminat Bedeli', 0),
            ])
        ws.append(output_row)

        # Hizalama: Tum sutunlar sola hizali
        for c in range(1, len(output_columns) + 1):
            ws.cell(idx, c).alignment = data_align_left

        # Number Formatting:
        # A(1), C(3), D(4), E(5), F(6), R(18) -> General
        for c in [1, 3, 4, 5, 6, 18]:
            if c <= len(output_columns):
                ws.cell(idx, c).number_format = 'General'

        # B(2) -> Sayı ('0')
        ws.cell(idx, 2).number_format = '0'

        # G(7)..Q(17) -> Virgül Stili (Güç, Kurulu, Aktif, Bedeller, KDV Matrahı, KDV, Toplam)
        for c in range(7, 18):
            ws.cell(idx, c).number_format = COMMA_STYLE_FORMAT

        if extended_output_enabled:
            for c in range(len(STANDARD_COLUMNS) + 1, len(output_columns) + 1):
                ws.cell(idx, c).number_format = COMMA_STYLE_FORMAT

    # Kaydetmeden önce tüm sütunları max karakter uzunluğu + 2 ile genişlet.
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if cell.value is not None and len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = max_length + 2
    
    try:
        wb.save(filepath)
        print(f"\n* Veri {filepath} konumuna kaydedildi ({len(data)} satır)")
    except PermissionError:
        print(f"\n[WARNING] {filepath} dosyası Excel'de açık olduğu için doğrudan kaydedilemedi!")
        alt_path = filepath.replace('.xlsx', '_Guncel.xlsx')
        wb.save(alt_path)
        print(f"* Veri alternatif olarak {alt_path} konumuna kaydedildi ({len(data)} satır)")
